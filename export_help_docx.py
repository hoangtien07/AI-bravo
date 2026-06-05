"""
export_help_docx.py — Xuất corpus help BRAVO 10 (data/raw/bravo10_help.jsonl) ra 1 file .docx
có cấu trúc theo chương/mục (breadcrumb path), kèm URL nguồn, để con người SOÁT LẠI.

Chạy:  .venv/Scripts/python.exe export_help_docx.py
Ra:    data/bravo10_help_huongdan.docx
"""
import json

from docx import Document
from docx.shared import Pt, RGBColor

import config

SRC = config.RAW_DIR / "bravo10_help.jsonl"
OUT = config.DATA_DIR / "bravo10_help_huongdan.docx"


def main():
    rows = [json.loads(l) for l in open(SRC, encoding="utf-8") if l.strip()]
    # Sắp theo path để gom chương/mục theo thứ tự cây.
    rows.sort(key=lambda r: (r.get("path") or r.get("title") or "").lower())

    doc = Document()
    doc.add_heading("Tài liệu Hướng dẫn sử dụng BRAVO 10", 0)
    p = doc.add_paragraph(
        f"Bản tổng hợp từ nội dung CÔNG KHAI help.bravo.com.vn đã thu thập "
        f"({len(rows)} trang). Dùng để SOÁT LẠI nội dung trước khi đưa vào AI. "
        f"Mỗi mục kèm đường dẫn nguồn để đối chiếu.")
    p.runs[0].italic = True

    last_chapter = None
    for r in rows:
        path = r.get("path") or r.get("title") or "(không rõ)"
        parts = [s.strip() for s in path.split(">") if s.strip()] or [path]
        chapter = parts[0]
        if chapter != last_chapter:           # đổi chương cấp 1 -> sang trang mới + H1
            doc.add_page_break()
            doc.add_heading(chapter, 1)
            last_chapter = chapter

        # Tiêu đề mục theo độ sâu (cấp 2..4); mục cấp 1 đã là heading chương ở trên.
        if len(parts) > 1:
            doc.add_heading(parts[-1], min(len(parts), 4))

        # URL nguồn (dòng nhỏ, xám).
        src = doc.add_paragraph()
        run = src.add_run(f"Nguồn: {r.get('url','')}")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        # Thân bài: tách theo dòng trống thành đoạn.
        text = (r.get("text") or "").strip()
        for para in text.split("\n\n"):
            para = para.strip()
            if para:
                doc.add_paragraph(para)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Đã xuất {len(rows)} mục -> {OUT}")


if __name__ == "__main__":
    main()
